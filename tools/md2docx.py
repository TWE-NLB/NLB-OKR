"""Markdown → Word (.docx) für die OKR-Dokumentation. Unterstützt Überschriften (#–####), Absätze, **fett**, *kursiv*,
[Links](url) und nackte URLs, Aufzählungen (- und eingerückt), Nummernlisten (1.), Tabellen, Hinweiskästen (>) und ---.
Rechts-nach-links (Hebräisch) mit --rtl. Aufruf: python3 md2docx.py ein.md aus.docx --lang de-DE --footer "Text" [--rtl]"""
import argparse
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ap = argparse.ArgumentParser()
ap.add_argument('src'); ap.add_argument('out')
ap.add_argument('--lang', default='de-DE'); ap.add_argument('--footer', default=''); ap.add_argument('--rtl', action='store_true')
ap.add_argument('--title', default='')
a = ap.parse_args()

ACCENT = RGBColor(0x1E, 0x5A, 0x7A)
INK2 = RGBColor(0x4E, 0x5A, 0x62)
FONT = 'Calibri'
FONT_CS = 'Arial'  # hebräische Schrift (Complex Script)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)


def set_fonts(rpr_parent, size=None):
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, 'get_or_add_rPr') else rpr_parent
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        fonts.set(qn(k), FONT)
    fonts.set(qn('w:cs'), FONT_CS)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); rpr.append(lang)
    lang.set(qn('w:val'), a.lang)
    if a.rtl:
        lang.set(qn('w:bidi'), a.lang)
    if size:
        szcs = rpr.find(qn('w:szCs'))
        if szcs is None:
            szcs = OxmlElement('w:szCs'); rpr.append(szcs)
        szcs.set(qn('w:val'), str(int(size * 2)))


# Grundformate
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(10.5)
set_fonts(st.element, 10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for name, size, color in (('Title', 22, ACCENT), ('Heading 1', 16, ACCENT), ('Heading 2', 13, ACCENT), ('Heading 3', 11.5, ACCENT)):
    s = doc.styles[name]
    s.font.name = FONT; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = color
    set_fonts(s.element, size)
    s.element.get_or_add_rPr().append(OxmlElement('w:bCs'))
    s.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
    if name == 'Title':
        # Titel ohne Linie und ohne Zierschrift
        pbdr = s.element.find(qn('w:pPr'))
        if pbdr is not None:
            for b in pbdr.findall(qn('w:pBdr')):
                pbdr.remove(b)
for name in ('List Bullet', 'List Bullet 2'):
    doc.styles[name].paragraph_format.space_after = Pt(3)


def bidi(p):
    if a.rtl:
        ppr = p._p.get_or_add_pPr()
        if ppr.find(qn('w:bidi')) is None:
            ppr.insert(0, OxmlElement('w:bidi'))


def add_hyperlink(p, text, url, bold=False):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1E5A7A'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    if bold:
        rpr.append(OxmlElement('w:b'))
    set_fonts(rpr)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)


INLINE = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[[^\]]+\]\([^)]+\)|https?://[^\s)<|]+|<br>)')


def add_inline(p, text, size=None, color=None, bold_all=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b = i = False; url = None; t = part
        if part == '<br>':
            p.add_run().add_break(); continue
        if part.startswith('***') and part.endswith('***'):
            t, b, i = part[3:-3], True, True
        elif part.startswith('**') and part.endswith('**'):
            t, b = part[2:-2], True
        elif re.fullmatch(r'\*[^*\s][^*]*?\*', part):
            t, i = part[1:-1], True
        elif part.startswith('[') and '](' in part:
            t, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part).groups()
        elif part.startswith('http'):
            t = url = part
        # Links innerhalb von Fett (z. B. **[1]**) nicht nötig; Fett mit Link-Inhalt wird als Text behandelt
        if url:
            add_hyperlink(p, t.replace('\\|', '|'), url, b or bold_all); continue
        r = p.add_run(t.replace('\\|', '|'))
        r.bold = b or bold_all or None
        r.italic = i or None
        # Fett/kursiv auch für hebräische Schrift (Complex Script)
        if b or bold_all:
            r._r.get_or_add_rPr().append(OxmlElement('w:bCs'))
        if i:
            r._r.get_or_add_rPr().append(OxmlElement('w:iCs'))
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        set_fonts(r._r, size)


def shade(cell_or_p_pr, fill):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    cell_or_p_pr.append(sh)


def add_box(lines):
    # Hinweiskasten: einspaltige Tabelle mit hellem Hintergrund und Rand
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles['Table Grid']
    if a.rtl:
        t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    c = t.rows[0].cells[0]
    shade(c._tc.get_or_add_tcPr(), 'FFF6E5')
    c.paragraphs[0].text = ''
    for k, line in enumerate(lines):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        bidi(p); add_inline(p, line)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(rows):
    n = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n)
    t.style = doc.styles['Table Grid']
    if a.rtl:
        t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    for ri, r in enumerate(rows):
        for ci in range(n):
            cell = t.rows[ri].cells[ci]
            p = cell.paragraphs[0]; bidi(p)
            p.paragraph_format.space_after = Pt(2)
            txt = r[ci] if ci < len(r) else ''
            if ri == 0:
                shade(cell._tc.get_or_add_tcPr(), '1E5A7A')
                add_inline(p, txt, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)
            else:
                add_inline(p, txt, size=9.5)
    # Kopfzeile auf jeder Seite wiederholen
    trpr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); trpr.append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def split_row(line):
    cells = re.split(r'(?<!\\)\|', line.strip().strip('|'))
    return [c.strip() for c in cells]


lines = open(a.src).read().split('\n')
i = 0
num = 0
first_heading = True
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1; num = 0 if not (i < len(lines) and re.match(r'^\d+\. ', lines[i])) else num; continue
    m = re.match(r'^(#{1,4}) (.*)$', line)
    if m:
        level = len(m.group(1))
        style = 'Title' if level == 1 else f'Heading {level - 1}'
        p = doc.add_paragraph(style=style); bidi(p); add_inline(p, m.group(2))
        i += 1; num = 0; continue
    if line.startswith('> '):
        box = []
        while i < len(lines) and lines[i].startswith('>'):
            box.append(lines[i][1:].strip()); i += 1
        add_box([b for b in box if b]); continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|(\s*-+\s*\|)+\s*$', lines[i]):
                rows.append(split_row(lines[i]))
            i += 1
        add_table(rows); continue
    if line.strip() == '---':
        p = doc.add_paragraph(); bidi(p)
        pbdr = OxmlElement('w:pBdr'); top = OxmlElement('w:top')
        for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '6'), ('w:color', 'BFC8CC')):
            top.set(qn(k), v)
        pbdr.append(top); p._p.get_or_add_pPr().append(pbdr)
        i += 1; continue
    m = re.match(r'^(\s*)- (.*)$', line)
    if m:
        p = doc.add_paragraph(style='List Bullet 2' if len(m.group(1)) >= 2 else 'List Bullet'); bidi(p)
        add_inline(p, m.group(2)); i += 1; continue
    m = re.match(r'^(\s*)(\d+)\. (.*)$', line)
    if m:
        num += 1
        p = doc.add_paragraph(); bidi(p)
        pf = p.paragraph_format; pf.left_indent = Cm(0.75); pf.first_line_indent = Cm(-0.75); pf.space_after = Pt(3)
        if a.rtl:  # Einzug am Zeilenanfang (rechts)
            pf.left_indent = None; pf.first_line_indent = None
            ind = OxmlElement('w:ind'); ind.set(qn('w:start'), '425'); ind.set(qn('w:hanging'), '425'); p._p.get_or_add_pPr().append(ind)
        r = p.add_run(f'{num}.\t'); set_fonts(r._r)
        add_inline(p, m.group(3)); i += 1; continue
    # normaler Absatz (Folgezeilen bis Leerzeile zusammenfassen)
    buf = [line.strip()]; i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r'^(#|>|\||- |\s+- |\d+\. |---)', lines[i]):
        buf.append(lines[i].strip()); i += 1
    p = doc.add_paragraph(); bidi(p); add_inline(p, ' '.join(buf))

# Fußzeile mit Seitenzahl
if a.footer:
    fp = sec.footer.paragraphs[0]; bidi(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(a.footer + ' '); r.font.size = Pt(8.5); r.font.color.rgb = INK2; set_fonts(r._r, 8.5)
    for kind in ('begin', None, 'end'):
        r = fp.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = INK2
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)

doc.core_properties.title = a.title or (lines[0].lstrip('# ').strip() if lines else '')
doc.core_properties.language = a.lang
doc.save(a.out)
print(a.out)
