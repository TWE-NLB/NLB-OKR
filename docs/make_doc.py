# -*- coding: utf-8 -*-
"""Erzeugt die Dokumentation der NLB OKR-Messung als .docx (python-docx)."""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
ACCENT = RGBColor(0x1E, 0x5A, 0x7A)
INK2 = RGBColor(0x4E, 0x5A, 0x62)

# ------------------------------------------------------------------ Literatur
REFS = [
    ("23477776", "Knapp CM, Proudlock FA, Gottlob I (2013). OKN asymmetry in human subjects: a literature review. Strabismus 21(1):37-49.", "10.3109/09273972.2012.762532", None),
    ("18235002", "Knapp CM, Gottlob I, McLean RJ, Proudlock FA (2008). Horizontal and vertical look and stare optokinetic nystagmus symmetry in healthy adult volunteers. Invest Ophthalmol Vis Sci 49(2):581-8.", "10.1167/iovs.07-0773", None),
    ("15864563", "Konen CS, Kleiser R, Seitz RJ, Bremmer F (2005). An fMRI study of optokinetic nystagmus and smooth-pursuit eye movements in humans. Exp Brain Res 165(2):203-16.", "10.1007/s00221-005-2289-7", None),
    ("20403339", "Kashou NH, Leguire LE, Roberts CJ et al. (2010). Instruction dependent activation during optokinetic nystagmus (OKN) stimulation: An FMRI study at 3T. Brain Res 1336:10-21.", "10.1016/j.brainres.2010.04.017", None),
    ("3384058", "van den Berg AV, Collewijn H (1988). Directional asymmetries of human optokinetic nystagmus. Exp Brain Res 70(3):597-604.", "10.1007/BF00247608", None),
    ("12939299", "Garbutt S, Han Y, Kumar AN et al. (2003). Vertical optokinetic nystagmus and saccades in normal human subjects. Invest Ophthalmol Vis Sci 44(9):3833-41.", "10.1167/iovs.03-0066", None),
    ("2792261", "Murasugi CM, Howard IP (1989). Up-down asymmetry in human vertical optokinetic nystagmus and afternystagmus: contributions of the central and peripheral retinae. Exp Brain Res 77(1):183-92.", "10.1007/BF00250580", None),
    ("26381838", "Waddington J, Harris CM (2015). Human optokinetic nystagmus and spatial frequency. J Vis 15(13):7.", "10.1167/15.13.7", None),
    ("15548816", "Valmaggia C, Rütsche A, Baumann A et al. (2004). Age related change of optokinetic nystagmus in healthy subjects: a study from infancy to senescence. Br J Ophthalmol 88(12):1577-81.", "10.1136/bjo.2004.044222", "PMC1772432"),
    ("8050519", "Paige GD (1994). Senescence of human visual-vestibular interactions: smooth pursuit, optokinetic, and vestibular control of eye movements with aging. Exp Brain Res 98(2):355-72.", "10.1007/BF00228423", None),
    ("16122840", "Kerber KA, Ishiyama GP, Baloh RW (2006). A longitudinal study of oculomotor function in normal older people. Neurobiol Aging 27(9):1346-53.", "10.1016/j.neurobiolaging.2005.07.009", None),
    ("26410670", "Seferlis F, Chimona TS, Papadakis CE et al. (2015). Age related changes in ocular motor testing in healthy subjects. J Vestib Res 25(2):57-66.", "10.3233/VES-150548", None),
    ("9712010", "Dieterich M, Bucher SF, Seelos KC, Brandt T (1998). Horizontal or vertical optokinetic stimulation activates visual motion-sensitive, ocular motor and vestibular cortex areas with right hemispheric dominance. An fMRI study. Brain 121(8):1479-95.", "10.1093/brain/121.8.1479", None),
    ("16080162", "Bense S, Janusch B, Schlindwein P et al. (2006). Direction-dependent visual cortex activation during horizontal optokinetic stimulation (fMRI study). Hum Brain Mapp 27(4):296-305.", "10.1002/hbm.20185", "PMC6871394"),
    ("16636788", "Bense S, Janusch B, Vucurevic G et al. (2006). Brainstem and cerebellar fMRI-activation during horizontal and vertical optokinetic stimulation. Exp Brain Res 174(2):312-23.", "10.1007/s00221-006-0464-0", None),
    ("28646240", "Ruehl RM, Hinkel C, Bauermann T, zu Eulenburg P (2017). Delineating function and connectivity of optokinetic hubs in the cerebellum and the brainstem. Brain Struct Funct 222(9):4163-85.", "10.1007/s00429-017-1461-8", None),
    ("2517458", "Hoffmann KP (1989). Control of the optokinetic reflex by the nucleus of the optic tract in primates. Prog Brain Res 80:173-82.", "10.1016/s0079-6123(08)62211-6", None),
    ("2323371", "Schiff D, Cohen B, Büttner-Ennever J, Matsuo V (1990). Effects of lesions of the nucleus of the optic tract on optokinetic nystagmus and after-nystagmus in the monkey. Exp Brain Res 79(2):225-39.", "10.1007/BF00608231", None),
    ("18718327", "Distler C, Hoffmann KP (2008). Private lines of cortical visual information to the nucleus of the optic tract and dorsolateral pontine nucleus. Prog Brain Res 171:363-8.", "10.1016/S0079-6123(08)00653-5", None),
    ("35074057", "Robinson DA (2022). Neurophysiology of the optokinetic system. Prog Brain Res 267(1):251-69.", "10.1016/bs.pbr.2021.10.012", None),
    ("3346723", "Maioli C (1988). Optokinetic nystagmus: modeling the velocity storage mechanism. J Neurosci 8(3):821-32.", "10.1523/JNEUROSCI.08-03-00821.1988", "PMC6569227"),
    ("19567703", "Baier B, Stoeter P, Dieterich M (2009). Anatomical correlates of ocular motor deficits in cerebellar lesions. Brain 132(8):2114-24.", "10.1093/brain/awp165", None),
    ("10372073", "Walker MF, Zee DS (1999). Directional abnormalities of vestibular and optokinetic responses in cerebellar disease. Ann N Y Acad Sci 871:205-20.", "10.1111/j.1749-6632.1999.tb09186.x", None),
    ("7066670", "Büttner-Ennever JA, Büttner U, Cohen B, Baumgartner G (1982). Vertical gaze paralysis and the rostral interstitial nucleus of the medial longitudinal fasciculus. Brain 105(1):125-49.", "10.1093/brain/105.1.125", None),
    ("2303833", "Bogousslavsky J, Miklossy J, Regli F, Janzer R (1990). Vertical gaze palsy and selective unilateral infarction of the rostral interstitial nucleus of the medial longitudinal fasciculus (riMLF). J Neurol Neurosurg Psychiatry 53(1):67-71.", "10.1136/jnnp.53.1.67", "PMC1014100"),
    ("11960796", "Sparks DL, Barton EJ, Gandhi NJ, Nelson J (2002). Studies of the role of the paramedian pontine reticular formation in the control of head-restrained and head-unrestrained gaze shifts. Ann N Y Acad Sci 956:85-98.", "10.1111/j.1749-6632.2002.tb02811.x", None),
    ("18718275", "Büttner-Ennever JA (2008). Mapping the oculomotor system. Prog Brain Res 171:3-11.", "10.1016/S0079-6123(08)00601-8", None),
    ("17314476", "Büttner-Ennever JA (2007). Anatomy of the oculomotor system. Dev Ophthalmol 40:1-14.", "10.1159/000100345", None),
    ("25145891", "Strupp M, Kremmyda O, Adamczyk C et al. (2014). Central ocular motor disorders, including gaze palsy and nystagmus. J Neurol 261 Suppl 2:S542-58.", "10.1007/s00415-014-7385-9", "PMC4141156"),
    ("15377682", "Garbutt S, Riley DE, Kumar AN et al. (2004). Abnormalities of optokinetic nystagmus in progressive supranuclear palsy. J Neurol Neurosurg Psychiatry 75(10):1386-94.", "10.1136/jnnp.2003.027367", "PMC1738726"),
    ("19514011", "Knapp CM, Gottlob I, McLean RJ et al. (2009). Vertical optokinetic nystagmus in Parkinson's disease. Mov Disord 24(10):1533-8.", "10.1002/mds.22634", None),
    ("666605", "Baloh RW, Yee RD, Honrubia V (1978). Internuclear ophthalmoplegia. II. Pursuit, optokinetic nystagmus, and vestibulo-ocular reflex. Arch Neurol 35(8):490-3.", "10.1001/archneur.1978.00500320010003", None),
    ("11677733", "Kaga K, Nakamura M, Furuya N et al. (2001). Analysis of optokinetic nystagmus in patients with lesions on the left unilateral parietal lobe or the entire left hemisphere. Acta Otolaryngol Suppl 545:166-9.", "10.1080/000164801750388360", None),
    ("9010000", "Heide W, Kurzidim K, Kömpf D (1996). Deficits of smooth pursuit eye movements after frontal and parietal lesions. Brain 119(6):1951-69.", "10.1093/brain/119.6.1951", None),
    ("12208005", "Doricchi F, Siegler I, Iaria G, Berthoz A (2002). Vestibulo-ocular and optokinetic impairments in left unilateral neglect. Neuropsychologia 40(12):2084-99.", "10.1016/s0028-3932(02)00049-0", None),
    ("11388418", "Todd L, King J, Darlington CL, Smith PF (2001). Optokinetic reflex dysfunction in multiple sclerosis. Neuroreport 12(7):1399-402.", "10.1097/00001756-200105250-00021", None),
    ("33367947", "Bertolini G, Romano F, Straumann D et al. (2021). Measuring optokinetic after-nystagmus: potential for detecting patients with signs of visual dependence following concussion. J Neurol 268(5):1747-61.", "10.1007/s00415-020-10359-8", "PMC8068696"),
    ("38607620", "Frattini D, Rosén N, Wibble T (2024). A proposed mechanism for visual vertigo: post-concussion patients have higher gain from visual input into subcortical gaze stabilization. Invest Ophthalmol Vis Sci 65(4):26.", "10.1167/iovs.65.4.26", "PMC11018265"),
    ("39258616", "Turuwhenua J, LinTun Z, Norouzifard M et al. (2024). Automated visual acuity estimation by optokinetic nystagmus using a stepped sweep stimulus. Ophthalmic Physiol Opt 44(7):1500-12.", "10.1111/opo.13391", None),
]
NUM = {pmid: i + 1 for i, (pmid, *_rest) in enumerate(REFS)}


def R(pmid, section="Results"):
    """Zitat im Text: [Nr., Abschnitt]."""
    return f"[{NUM[pmid]}, {section}]"


# ------------------------------------------------------------------ Dokument
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

st = doc.styles
st["Normal"].font.name = "Calibri"
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st["Normal"].font.size = Pt(10.5)
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.15
for name, size, color in (("Title", 24, ACCENT), ("Heading 1", 16, ACCENT), ("Heading 2", 13, ACCENT), ("Heading 3", 11.5, RGBColor(0x13, 0x1A, 0x1E))):
    s = st[name]
    s.font.name = "Calibri"
    s.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), "Calibri")
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True



LIST = {"id": None}


def H(text, level=1):
    LIST["id"] = None
    return doc.add_heading(text, level=level)


def new_num_id():
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style = doc.styles["List Number"]
    num_id = style.element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(num_id).abstractNumId.val
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


def fix(text):
    return text.replace("][", "; ") if isinstance(text, str) else text


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def P(text="", bold_lead=None, style=None, italic=False, size=None, color=None, align=None):
    text = fix(text)
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    if text:
        r = p.add_run(text); r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    if align: p.alignment = align
    return p


def B(text, bold_lead=None, level=0):
    text = fix(text)
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def N(text, bold_lead=None):
    text = fix(text)
    p = doc.add_paragraph(style="List Number")
    if LIST["id"] is None:
        LIST["id"] = new_num_id()
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = LIST["id"]
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def T(header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1E5A7A")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            if isinstance(val, tuple):
                r = para.add_run(val[0]); r.bold = True; r.font.size = Pt(9.5)
                r2 = para.add_run(val[1]); r2.font.size = Pt(9.5)
            else:
                r = para.add_run(fix(str(val))); r.font.size = Pt(9.5)
            para.paragraph_format.space_after = Pt(0)
            if ri % 2 == 1:
                shade(cells[i], "EEF3F5")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # Kopfzeile auf jeder Seite wiederholen
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def note(text, lead="Hinweis: "):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, "FFF4DC")
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(lead); r.bold = True; r.font.size = Pt(10)
    r2 = p.add_run(text); r2.font.size = Pt(10)
    c.width = Cm(16.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def measure(title, definition, calc, normal, meaning, anatomy, changes, compare):
    H(title, level=2)
    P(definition, "Definition: ")
    P(calc, "So berechnet die App: ")
    H("Normalwerte und Einflussfaktoren", level=3)
    for lead, txt in normal: B(txt, lead)
    H("Bedeutung", level=3)
    for lead, txt in meaning: B(txt, lead)
    H("Beteiligte Hirnbereiche und Augenmuskeln", level=3)
    for lead, txt in anatomy: B(txt, lead)
    if changes:
        H("Typische Veränderungen in Studien", level=3)
        for lead, txt in changes: B(txt, lead)
    H("Vergleich von zwei oder mehr Messungen", level=3)
    for lead, txt in compare: B(txt, lead)


def toc():
    p = doc.add_paragraph()
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = 'TOC \\o "1-2" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Inhaltsverzeichnis: In Word mit Rechtsklick → „Feld aktualisieren“ erzeugen."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3): r._r.append(el)
    # Word fragt beim Öffnen, ob Felder aktualisiert werden sollen
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)


def footer():
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("NLB OKR-Messung · Dokumentation · Seite ")
        r.font.size = Pt(8.5); r.font.color.rgb = INK2
        run = p.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("separate", None), (None, None), ("end", None)):
            if typ:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ); run._r.append(el)
            elif txt:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; run._r.append(el)
        run.font.size = Pt(8.5); run.font.color.rgb = INK2


# ================================================================== Inhalt
doc.add_paragraph("NLB OKR-Messung", style="Title")
P("Dokumentation der App und Erklärung der Messwerte", size=14, color=INK2)
P("Stand: 14. Oktober 2026 · App-Version 1 · https://twe-nlb.github.io/NLB-OKR/", size=9.5, color=INK2)

note("Die App ist eine Demo und kein Medizinprodukt. Sie stellt keine Diagnose. Die zitierten Studien wurden mit Labormethoden "
     "(Magnetspulen, Infrarot-Okulographie, große Reizfelder) durchgeführt. Die Befunde zeigen, welche Hirnbereiche einen Messwert "
     "beeinflussen können. Sie lassen sich nicht direkt auf eine einzelne Messung mit der iPad-Kamera übertragen. "
     "Auffällige Werte gehören ärztlich (Neurologie, Neuro-Ophthalmologie, HNO) abgeklärt.", "Wichtig: ")

P("Literaturverweise stehen im Text in eckigen Klammern: [Nummer, Abschnitt]. Die Nummer verweist auf das Literaturverzeichnis "
  "am Ende. Der Abschnitt bezeichnet den Teil der Studienzusammenfassung (Abstract), aus dem die Aussage stammt, zum Beispiel "
  "„Results“ oder „Conclusions“. Bei Abstracts ohne Gliederung steht „Abstract“. Ausgewertet wurden die PubMed-Abstracts. "
  "Die Volltexte wurden nicht gelesen.", "Zitierweise: ")

H("Inhalt", level=1)
toc()
doc.add_page_break()

# ---------------------------------------------------------------- 1
H("1 Der optokinetische Reflex in Kürze", level=1)
P("Bewegt sich ein großes Muster durch das Gesichtsfeld, folgen die Augen ihm automatisch. Das ist der optokinetische Reflex (OKR). "
  "Die Augenbewegung, die dabei entsteht, heißt optokinetischer Nystagmus (OKN). Er besteht aus zwei Teilen "
  + R("23477776", "Abstract") + ":")
B("Das Auge folgt dem Muster gleichmäßig in Laufrichtung. Diese Phase sagt am meisten über die Bewegungsverarbeitung aus.", "Langsame Phase: ")
B("Ein kurzer Rücksprung (eine Sakkade) gegen die Laufrichtung setzt das Auge zurück.", "Schnelle Phase: ")
P("Im zeitlichen Verlauf ergibt das einen Sägezahn. Dessen Rhythmus und Höhe schwanken auch bei Gesunden von Schlag zu Schlag deutlich "
  + R("26381838", "Abstract") + ".")
H("Look-OKN und Stare-OKN", level=2)
P("Es gibt zwei Formen des OKN, und welche entsteht, hängt von der Anweisung an die Testperson ab:")
B("Die Person verfolgt bewusst einzelne Streifen. Die langsamen Phasen sind lang und die Rücksprünge groß und selten.", "Look-OKN: ")
B("Die Person schaut passiv geradeaus. Das ist eher reflexhaft, mit kürzeren langsamen Phasen und niedrigerem Gain "
  + R("18235002", "Purpose") + ".", "Stare-OKN: ")
P("Look-OKN aktiviert deutlich mehr Hirnrinde als Stare-OKN " + R("20403339", "Abstract") + R("15864563", "Abstract")[0:0] + ". "
  "Stare-OKN wird überwiegend von Hirnstamm und Kleinhirn gesteuert " + R("28646240", "Abstract") + ". "
  "Die App gibt die Anweisung „Entspannt auf die Streifen schauen, ihnen folgen ist erlaubt“. Das ergibt meist eine Mischform mit Look-Anteil. "
  "Für vergleichbare Messungen muss die Anweisung deshalb immer gleich sein.")

# ---------------------------------------------------------------- 2
H("2 Die App", level=1)
H("2.1 Technik", level=2)
B("Die App ist eine Webseite. Sie läuft in Safari auf iPad oder iPhone, am besten als App auf dem Home-Bildschirm.", "Web-App: ")
B("Die Frontkamera liefert etwa 30 Bilder pro Sekunde, auf manchen Geräten bis 60.", "Kamera: ")
B("Das Programm „MediaPipe Face Landmarker“ findet in jedem Kamerabild 478 Gesichtspunkte, darunter die Iris-Mittelpunkte und die Augenwinkel.", "Augenerkennung: ")
B("Die App rechnet die Lage der Iris zwischen den Augenwinkeln aus und teilt sie durch die Augenbreite. "
  "Dadurch bleibt der Wert weitgehend gleich, wenn sich Kopf oder Gerät im Bild verschieben. "
  "Die App bildet den Mittelwert beider Augen.", "Augenposition: ")
B("Kameragestützte, automatisch ausgewertete OKN-Messung ist grundsätzlich machbar. Eine Studie bestimmte damit die Sehschärfe ähnlich gut wie mit der Sehtafel " + R("39258616", "Conclusions") + ".", "Vorbild: ")
B("Alle Messungen bleiben auf dem Gerät. Die App speichert kein Video. Gesichert wird über den Export als Datei.", "Daten: ")

H("2.2 Ablauf einer Messung", level=2)
N("Testperson unter einem Pseudonym anlegen.", "Person: ")
N("Haltung (Ständer oder Hand), Abstand, Gerät, Richtungen, Geschwindigkeiten, Blockdauer, Streifenbreite und Reihenfolge festlegen.", "Einstellungen: ")
N("Die App prüft Kamera, Gesicht, Augen, Licht, Position und Ruhe. Alles muss grün sein, und der gemessene Abstand muss bestätigt werden.", "Setup-Check: ")
N("Ein Punkt erscheint nacheinander Mitte → links → Mitte → rechts → Mitte → oben → Mitte → unten → Mitte, je 1,8 s lang. "
  "Die Messung startet nur, wenn alle Prüfkriterien erfüllt sind (Abschnitt 2.4).", "Kalibrierung: ")
N("Pro Block zeigt die App 3 s ein Fixierkreuz mit Richtungspfeil. Danach läuft das Streifenband, standardmäßig 10 s lang.", "Messung: ")
N("Die App zeigt Kennwerte, eine Tabelle aller Blöcke und die Augenkurve pro Block.", "Ergebnis: ")
N("Bei zwei oder mehr angehakten Messungen zeigt die App Tabelle, Gain-Diagramme und Asymmetrie-Diagramm.", "Vergleich: ")
N("JSON (vollständig, wieder importierbar) und CSV (für Excel) über das Teilen-Menü: Dateien, AirDrop oder SeaFile.", "Export: ")

H("2.3 Einstellungen und Standardwerte", level=2)
T(["Einstellung", "Standard", "Bedeutung"], [
    ["Haltung", "Ständer", "Mit „In der Hand“ gelten großzügigere Grenzen für Bewegung. Die Einstellung wird mit der Messung gespeichert."],
    ["Abstand Auge–Bildschirm", "40 cm", "Rechnet Bildschirmpunkte in Sehwinkel (Grad) um. Bitte messen, nicht schätzen."],
    ["Gerät (Punkte pro Zoll)", "iPad 132, iPad mini 163, iPhone 153", "Physische Größe der Bildschirmpunkte. Nötig für die Umrechnung in Grad."],
    ["Richtungen", "← → ↑ ↓", "Laufrichtung der Streifen. Senkrechte Streifen bei waagerechter Bewegung, waagerechte Streifen bei senkrechter."],
    ["Geschwindigkeiten", "10, 20, 30 °/s", "Streifengeschwindigkeit im Sehwinkel."],
    ["Dauer je Block", "10 s", "Die erste Sekunde wird verworfen, weil der Reflex dann noch anläuft."],
    ["Streifenbreite", "4° pro Paar (0,25 Perioden/°)", "Breite eines schwarzen plus eines weißen Streifens."],
    ["Reihenfolge", "je Geschwindigkeit alle Richtungen", "Verteilt Ermüdung gleichmäßig auf die Richtungen."],
], [4.2, 4.2, 8.2])

H("2.4 Setup-Check und Kalibrierung", level=2)
P("Der Setup-Check läuft vor der Kalibrierung. Rot blockiert den Start, Gelb ist erlaubt.")
T(["Prüfpunkt", "Grün, wenn …"], [
    ["Kamera", "mindestens 20 Bilder pro Sekunde (Gelb ab 12)"],
    ["Gesicht erkannt", "in mindestens 90 % der Bilder der letzten 1,5 s"],
    ["Augen offen, Iris sichtbar", "Lidspalte mindestens 18 % der Augenbreite"],
    ["Licht", "mittlere Bildhelligkeit zwischen 27 % und 90 %"],
    ["Gesicht mittig", "Augenmitte im mittleren Bereich des Kamerabildes"],
    ["Kopf ruhig / Gerät ruhig", "Schwankung unter 3 % des Augenabstands (in der Hand: 8 %)"],
    ["Abstand", "Grobe Schätzung als Information; der gemessene Abstand muss bestätigt werden"],
], [5.0, 11.6])
P("Kalibrierung: Die Kriterien unten müssen alle erfüllt sein, bevor das Streifenband startet.")
T(["Kriterium", "Grenze", "Wozu"], [
    ["Erkennungsrate je Punkt", "≥ 60 %", "Genug gültige Bilder pro Blickziel"],
    ["Waagerechte Anpassung (R²)", "≥ 0,8", "Augenmerkmal und Blickwinkel hängen linear zusammen"],
    ["Waagerechte Trennung", "≥ 4 × Rauschen", "Links und rechts klar unterscheidbar"],
    ["Senkrechte Anpassung (R²)", "≥ 0,7", "Senkrecht ist die Iris-Bewegung kleiner und schwerer zu messen"],
    ["Senkrechte Trennung", "≥ 3 × Rauschen", "Oben und unten klar unterscheidbar"],
    ["Wiederholbarkeit der Mitte", "Streuung ≤ 1,5° (waagerecht) / 2,5° (senkrecht)", "Fünfmaliges Fixieren der Mitte stimmt überein"],
    ["Kopfbewegung", "≤ 25 % des Augenabstands (Hand: 60 %)", "Kopf und Gerät standen ruhig"],
], [4.6, 4.8, 7.2])

H("2.5 Qualitätskontrolle während der Messung", level=2)
B("Lidschlag: Die Lidspalte ist kleiner als 55 % des Kalibrierwerts. Solche Bilder gelten als ungültig.", "Ungültige Bilder: ")
B("Gerätedrehung über 20 °/s (Gyroskop) oder Gesichtsbewegung im Kamerabild über 1 Augenabstand pro Sekunde. "
  "Diese Bilder fließen nicht in den Gain ein.", "Wackeln: ")
B("Kommt das Gesicht näher, wird der Sehwinkel der Streifen größer. Die App rechnet die Streifengeschwindigkeit deshalb über die Gesichtsgröße auf den tatsächlichen Abstand um.", "Abstandskorrektur: ")
B("„gut“: mindestens 70 % gültige Bilder und höchstens 20 % Wackeln. „mäßig“: mindestens 50 % und höchstens 40 %. "
  "Sonst „ungültig“, dann schlägt die App vor, den Block zu wiederholen.", "Blockqualität: ")

H("2.6 Grenzen der Messgenauigkeit", level=2)
B("Laborverfahren (Magnetspule, Infrarot-Okulographie) tasten deutlich schneller ab. Die App misst mit etwa 30 Bildern pro Sekunde. "
  "Die langsame Phase und damit der Gain sind gut erfassbar. Schnelle Phasen dauern nur 20–50 ms und werden deshalb nur grob erfasst.", "Bildrate: ")
B("Das iPad deckt bei 40 cm etwa 30° × 22° des Gesichtsfeldes ab. Studien nutzen oft 60–70° " + R("12939299", "Methods") + ". "
  "Kleinere Reizfelder senken den Gain. Wird nur die Gesichtsfeldmitte gereizt, verschwindet die Oben-unten-Asymmetrie "
  + R("2792261", "Abstract") + ".", "Reizfeldgröße: ")
B("Die App bildet den Mittelwert beider Augen. Unterschiede zwischen den Augen, etwa bei einer internukleären Ophthalmoplegie, kann sie nicht erkennen.", "Ein Wert für beide Augen: ")
B("Für die App gibt es noch keine Studien zur Messwiederholbarkeit. Wie stark die Werte von Messung zu Messung schwanken, muss man daher pro Person selbst ermitteln (Abschnitt 5).", "Keine Validierungsdaten: ")

# ---------------------------------------------------------------- 3
H("3 Wie das Gehirn den OKR erzeugt", level=1)
P("Die Kette vom Auge bis zum Augenmuskel, vereinfacht:")
N("Die Netzhaut meldet Bildbewegung, also Bildverschiebung auf der Netzhaut (retinal slip).", "Netzhaut: ")
N("Der Nucleus des Tractus opticus (NOT) und der dorsale Terminalkern (DTN) im Prätektum, am Übergang von Thalamus zu Mittelhirn, "
  "enthalten richtungsempfindliche Nervenzellen. Der linke NOT/DTN bevorzugt Bewegung nach links und umgekehrt "
  + R("2517458", "Abstract") + ". Beim Menschen ist diese Region bei waagerechtem und senkrechtem OKN aktiv " + R("16636788", "Abstract") + ".",
  "Prätektum (NOT/DTN): ")
N("MT/V5 und MST, im Übergang von Schläfen- zu Hinterhauptslappen, verarbeiten Bewegung. Sie senden Informationen an NOT/DTN und an die Brückenkerne "
  + R("18718327", "Abstract") + ". Bei OKN sind visuelle Bewegungsareale, parietale und frontale Augenfelder aktiv, "
  "mit Betonung der rechten Hirnhälfte " + R("9712010", "Abstract") + ".", "Großhirnrinde: ")
N("Die Vestibulariskerne fassen visuelle und Gleichgewichtssignale zur Schätzung der Eigendrehung zusammen " + R("35074057", "Abstract") + ". "
  "Ein Speicher (velocity storage) hält die Augengeschwindigkeit aufrecht. Er erzeugt auch den Nachnystagmus im Dunkeln "
  + R("3346723", "Abstract") + R("2323371", "Abstract")[0:0] + ".", "Vestibulariskerne und Geschwindigkeitsspeicher: ")
N("Flocculus, Uvula und okulomotorischer Vermis sind beim OKN aktiv " + R("16636788", "Abstract") + R("28646240", "Abstract") + ". "
  "Schäden an Uvula und Pyramis gehen mit einem verminderten Gain der langsamen Phase einher " + R("19567703", "Abstract") + ".", "Kleinhirn: ")
N("Die schnellen Phasen entstehen in zwei Zentren " + R("18718275", "Abstract") + ":", "Sakkaden-Generatoren: ")
B("waagerecht in der paramedianen pontinen Formatio reticularis (PPRF) der Brücke " + R("11960796", "Abstract") + ",", level=1)
B("senkrecht im rostralen interstitiellen Kern des medialen Längsbündels (riMLF) im Mittelhirn " + R("7066670", "Abstract") + ".", level=1)
P("Die fMRT zeigt beim waagerechten OKN eine Aktivierung der PPRF und beim senkrechten OKN des riMLF " + R("16636788", "Abstract") + ".")
N("Kerne III (N. oculomotorius), IV (N. trochlearis) und VI (N. abducens) steuern die äußeren Augenmuskeln "
  + R("17314476", "Abstract") + ".", "Augenmuskelkerne: ")

H("3.1 Augenmuskeln je Bewegungsrichtung", level=2)
P("Die folgende Zuordnung ist anatomisches Lehrbuchwissen. Die zitierten Arbeiten belegen die Kerne und Verbindungen.")
T(["Augenbewegung", "Muskeln", "Nerv / Kern", "Vermittelt durch"], [
    ["nach rechts (beide Augen)", "rechts M. rectus lateralis + links M. rectus medialis", "rechts VI, links III",
     "Abducens-Interneurone über das mediale Längsbündel (MLF) " + R("25145891", "Abstract")],
    ["nach links (beide Augen)", "links M. rectus lateralis + rechts M. rectus medialis", "links VI, rechts III", "wie oben, gespiegelt"],
    ["nach oben", "M. rectus superior + M. obliquus inferior", "III", "riMLF, Nucleus interstitialis Cajal " + R("25145891", "Abstract")],
    ["nach unten", "M. rectus inferior + M. obliquus superior", "III und IV", "riMLF " + R("2303833", "Abstract")],
], [3.4, 5.0, 2.8, 5.4])
P("Laufen die Streifen nach rechts, bewegen sich beide Augen in der langsamen Phase nach rechts. Der Rücksprung geht nach links. "
  "Waagerechte Sakkaden werden jeweils von der PPRF auf der Seite erzeugt, in deren Richtung sie gehen " + R("11960796", "Abstract") + ".")

# ---------------------------------------------------------------- 4
H("4 Die Messwerte", level=1)
P("Für jeden Messwert folgen: Definition, Berechnung in der App, Normalwerte, Bedeutung, beteiligte Hirnbereiche und Augenmuskeln, "
  "typische Veränderungen in Studien und Hinweise zum Vergleich mehrerer Messungen.")

measure(
    "4.1 Gain (Verstärkung der langsamen Phase)",
    "Augengeschwindigkeit in der langsamen Phase geteilt durch die Streifengeschwindigkeit. 1,0 heißt: Das Auge läuft genauso schnell wie die Streifen. 0,5 heißt: halb so schnell.",
    "Augenposition in Grad (über die Kalibrierung) → Geschwindigkeit von Bild zu Bild → schnelle Phasen und Wackeln entfernen → "
    "Median der verbleibenden Geschwindigkeiten ÷ abstandskorrigierte Streifengeschwindigkeit. Für jeden Block, also jede Richtung und jede Geschwindigkeit.",
    [
        ("Gesunde Erwachsene: ", "Der Gain lag immer unter 0,85 und nahm mit steigender Geschwindigkeit ab " + R("5" if False else "3384058", "Abstract") + "."),
        ("Alter: ", "Der Gain steigt in den ersten Lebensjahren, bleibt bis etwa 50 Jahre stabil und sinkt danach um 6–18 % " + R("15548816", "Results") + ". "
         "Auch eine Studie mit 250 Gesunden von 18 bis 70 Jahren fand eine Verschlechterung mit dem Alter " + R("26410670", "Abstract") + ". "
         "Bei Folgebewegung auf ein großes Streifenmuster sanken Gain und Präzision mit zunehmendem Alter fortschreitend " + R("8050519", "Abstract") + "."),
        ("Reiz: ", "Der Gain ist bei niedriger Geschwindigkeit und groben Streifen am höchsten " + R("12939299", "Results") + R("12939299", "Conclusions")[0:0] + ". "
         "Eine andere Studie fand dagegen höhere Augengeschwindigkeiten bei feineren Streifen " + R("26381838", "Abstract") + ". "
         "Deshalb die Streifenbreite für Vergleiche nie ändern."),
        ("Beidäugig: ", "Beidäugiges Sehen ergibt einen höheren Gain als einäugiges " + R("3384058", "Abstract") + "."),
        ("Anweisung: ", "Look-OKN hat einen höheren Gain als Stare-OKN " + R("18235002", "Purpose") + "."),
    ],
    [
        ("Allgemein: ", "Der Gain zeigt, wie gut das Gehirn Bildbewegung erkennt und in eine gleich schnelle Augenbewegung umsetzt. Ziel ist ein stabiles Netzhautbild."),
        ("Bei der App: ", "Wegen des kleineren Reizfelds sind etwas niedrigere Werte als in Laborstudien zu erwarten (Abschnitt 2.6)."),
    ],
    [
        ("Eingang: ", "Netzhaut → NOT/DTN " + R("2517458", "Abstract") + " und kortikale Bewegungsareale MT/MST " + R("18718327", "Abstract") + "."),
        ("Verarbeitung: ", "Vestibulariskerne und Geschwindigkeitsspeicher " + R("35074057", "Abstract") + ", Kleinhirn (Uvula, Pyramis, Flocculus) " + R("19567703", "Abstract") + R("16636788", "Abstract")[0:0] + "."),
        ("Ausführung: ", "waagerecht M. rectus lateralis und medialis (III, VI), senkrecht M. rectus superior/inferior und Mm. obliqui (III, IV)."),
    ],
    [
        ("Kleinhirn: ", "Patienten mit Schäden an Uvula und Pyramis hatten einen verminderten Gain der langsamen OKN-Phase " + R("19567703", "Abstract") + "."),
        ("Kleinhirndegeneration: ", "Waagerechte optokinetische Reize lösten teils eine zusätzliche Drehbewegung des Auges um die Blickachse aus. Die App misst diese Komponente nicht " + R("10372073", "Abstract") + "."),
        ("Multiple Sklerose: ", "Die Geschwindigkeit der langsamen Phase in der Anstiegszeit war vermindert " + R("11388418", "Abstract") + "."),
        ("PSP: ", "Der senkrechte Gain war bei progressiver supranukleärer Blickparese reduziert, bei Parkinson meist normal " + R("15377682", "Results") + "."),
        ("Alter: ", "Bei über 75-Jährigen nahm der OKN-Gain über 9–12 Jahre ab. Der Rückgang hing mit Gang- und Gleichgewichtswerten zusammen " + R("16122840", "Abstract") + "."),
    ],
    [
        ("Gleiche Person: ", "Sinkt der Gain in allen Richtungen gleichmäßig, liegt das am häufigsten an Aufmerksamkeit, Müdigkeit, Licht oder einem anderen Abstand. "
         "Erst wenn er sich bei Wiederholung bestätigt, ist an eine allgemeine Veränderung des langsamen Folgesystems zu denken (Kleinhirn, Alter)."),
        ("Verschiedene Personen: ", "Nur Personen ähnlichen Alters vergleichen " + R("15548816", "Results") + ", mit gleichem Gerät, Abstand und gleicher Haltung."),
        ("Darstellung in der App: ", "Die Tabelle zeigt den Gain je Richtung und Geschwindigkeit sowie die Abweichung zur ältesten Messung absolut und in %. "
         "Das Diagramm „Gain je Richtung“ zeigt eine Linie pro Messung."),
    ],
)

measure(
    "4.2 Gain-Verlauf über die Geschwindigkeit",
    "Wie stark der Gain von 10 über 20 auf 30 °/s abfällt, also die Steigung der Linien im Diagramm „Gain je Richtung“.",
    "Kein eigener Wert. Die App zeigt den Verlauf im Ergebnis (Blocktabelle) und im Vergleich (vier Diagramme).",
    [
        ("Normal: ", "Der Gain nimmt mit steigender Geschwindigkeit ab " + R("3384058", "Abstract") + R("12939299", "Results")[0:0] + ". "
         "Senkrecht fällt er steiler ab als waagerecht " + R("3384058", "Abstract") + "."),
        ("Ursache: ", "Ab einer gewissen Geschwindigkeit kann das Folgesystem nicht mehr mithalten. Der Gain „sättigt“."),
    ],
    [
        ("Allgemein: ", "Der Verlauf zeigt die Leistungsreserve des Folgesystems. Ein stärkerer Abfall als zuvor oder als bei Vergleichspersonen deutet darauf hin, dass schnelle Reize schlechter verarbeitet werden."),
    ],
    [
        ("Beteiligt: ", "Vor allem das schnelle, direkte Folgesystem über kortikale Bewegungsareale, Brückenkerne und Flocculus " + R("2323371", "Abstract") + "."),
        ("Messung beim Menschen: ", "Die fMRT fand geschwindigkeitsabhängige Aktivierung in der Brücke (NRTP, PPRF), in Uvula, Flocculus und Kleinhirntonsillen " + R("28646240", "Abstract") + "."),
    ],
    [],
    [
        ("Gleiche Person: ", "Fällt der Gain nur bei 30 °/s stärker ab als früher, ist das eher eine Veränderung der Folgegeschwindigkeit als der Bewegungswahrnehmung insgesamt."),
        ("Lesehilfe: ", "Parallel verschobene Linien heißen: Der Gain ist insgesamt verändert. Verschieden steile Linien heißen: Die Geschwindigkeitsabhängigkeit ist verändert."),
    ],
)

measure(
    "4.3 Geschwindigkeit der langsamen Phase",
    "Augengeschwindigkeit in Grad pro Sekunde während der Folgebewegung (in der CSV: „Langsame Phase °/s“).",
    "Median der Bild-zu-Bild-Geschwindigkeit, ohne Rücksprünge und Wackeln. Der Gain ist diese Größe geteilt durch die Streifengeschwindigkeit.",
    [
        ("Normal: ", "Bei 10 °/s liegt sie knapp unter 10 °/s. Bei höheren Geschwindigkeiten bleibt sie zunehmend dahinter zurück " + R("3384058", "Abstract") + "."),
    ],
    [
        ("Allgemein: ", "Der absolute Wert ist bei Studien mit unterschiedlichen Streifengeschwindigkeiten leichter vergleichbar. Für Vergleiche innerhalb der App ist der Gain besser geeignet."),
    ],
    [("Beteiligt: ", "wie beim Gain (Abschnitt 4.1).")],
    [
        ("Gehirnerschütterung: ", "Nach Gehirnerschütterung waren die langsamen Phasen bei Drehreizen um die Blickachse schneller "
         "und der visuelle Gain höher " + R("38607620", "Results") + ". Die App misst diese Reizform nicht."),
    ],
    [("Vergleich: ", "Nur bei gleicher Streifengeschwindigkeit sinnvoll, sonst den Gain verwenden.")],
)

measure(
    "4.4 Schlagfrequenz",
    "Anzahl der Rücksprünge (schnellen Phasen) pro Sekunde.",
    "Rücksprünge werden erkannt, wenn das Auge schneller als max(25 °/s, 1,2 × Streifengeschwindigkeit) gegen die Laufrichtung springt "
    "und dabei mindestens 0,5° zurücklegt. Anzahl ÷ ausgewertete Zeit.",
    [
        ("Typisch: ", "Stare-OKN hat häufige kleine Rücksprünge, Look-OKN seltene große " + R("15864563", "Abstract") + "."),
        ("Zusammenhang: ", "Frequenz und Amplitude hängen von der Geschwindigkeit der langsamen Phase ab und schwanken stark von Schlag zu Schlag " + R("26381838", "Abstract") + "."),
    ],
    [
        ("Allgemein: ", "Die Frequenz zeigt vor allem, wie die Person schaut (Strategie, Aufmerksamkeit), und wie gut die schnellen Phasen ausgelöst werden."),
        ("Grenze der App: ", "Bei etwa 30 Bildern pro Sekunde können sehr kleine, schnell aufeinanderfolgende Rücksprünge verschmelzen oder fehlen."),
    ],
    [
        ("Waagerechte Rücksprünge: ", "Burst-Neurone der PPRF in der Brücke " + R("11960796", "Abstract") + R("18718275", "Abstract") + "."),
        ("Senkrechte Rücksprünge: ", "riMLF im Mittelhirn. Wird es beidseits zerstört, sind schnelle senkrechte Augenbewegungen gestört " + R("7066670", "Abstract") + "."),
        ("Ausführung: ", "dieselben Muskeln wie in der langsamen Phase, aber in Gegenrichtung (Abschnitt 3.1)."),
    ],
    [
        ("Neglect: ", "Bei Patienten mit Neglect der linken Seite nach Schäden der rechten Hirnhälfte waren Geschwindigkeit, Amplitude und Frequenz der Phasen mit langsamer Bewegung nach rechts vermindert " + R("12208005", "Abstract") + "."),
        ("Multiple Sklerose: ", "Die Schlagfrequenz war in einer Studie nicht verändert, wohl aber Geschwindigkeit und Amplitude " + R("11388418", "Abstract") + "."),
    ],
    [
        ("Gleiche Person: ", "Ändert sich nur die Frequenz und der Gain bleibt gleich, wurde meist anders geschaut (mehr „Look“ oder mehr „Stare“). Das ist kein Hinweis auf eine veränderte Hirnfunktion."),
        ("Bedeutsamer: ", "Bleiben Rücksprünge nur in einer Richtung aus oder werden seltener, während der Gain erhalten ist. Dann ist das Sakkadensystem dieser Richtung betroffen, nicht das Folgesystem."),
    ],
)

measure(
    "4.5 Amplitude",
    "Größe der Rücksprünge in Grad, als Median aller erkannten Rücksprünge eines Blocks.",
    "Größter Positionsabfall innerhalb eines erkannten Rücksprungs, in Grad über die Kalibrierung.",
    [
        ("Normal: ", "Über 95 % der senkrechten Rücksprünge sind kleiner als 10° " + R("12939299", "Results") + "."),
        ("Anweisung: ", "Look-OKN hat große, Stare-OKN kleine Amplituden " + R("15864563", "Abstract") + "."),
    ],
    [
        ("Allgemein: ", "Die Amplitude zeigt, wie weit das Auge dem Muster folgt, bevor es zurückspringt. Zusammen mit der Frequenz beschreibt sie die Form des Sägezahns."),
    ],
    [
        ("Beteiligt: ", "Sakkaden-Generatoren (PPRF, riMLF) und Augenmuskelkerne " + R("18718275", "Abstract") + "."),
        ("Schlagfeld: ", "Bei Nystagmus verschiebt sich die mittlere Blickposition um etwa 2,6° in Richtung der schnellen Phasen. Das verändert auch die Aktivierung der Sehrinde " + R("16080162", "Abstract") + "."),
    ],
    [
        ("Hirnhälfte: ", "Bei Schäden der gesamten linken Hirnhälfte hatte der OKN zur Gegenseite eine niedrige Amplitude und eine verminderte Geschwindigkeit der langsamen Phase " + R("11677733", "Abstract") + "."),
        ("Multiple Sklerose: ", "Die Amplitude war vermindert " + R("11388418", "Abstract") + "."),
        ("PSP: ", "Die senkrechten Rücksprünge waren verlangsamt und mit kleinen waagerechten Sakkaden kombiniert " + R("15377682", "Results") + "."),
    ],
    [
        ("Gleiche Person: ", "Wie bei der Frequenz zuerst an eine andere Blickstrategie denken. Bedeutsam ist eine Richtung, in der die Amplitude allein abnimmt, zusammen mit einem Gain-Unterschied."),
    ],
)

measure(
    "4.6 Asymmetrie rechts/links",
    "(Gain rechts − Gain links) ÷ Mittelwert beider × 100 %. Positiv: Nach rechts laufende Streifen werden besser verfolgt.",
    "Mittelwert der gültigen Blöcke je Richtung über alle Geschwindigkeiten, dann die Formel oben.",
    [
        ("Normal: ", "Waagerecht gibt es bei Gesunden keine Gruppenasymmetrie " + R("3384058", "Abstract") + R("18235002", "Results") + ". "
         "Nach einer Übersichtsarbeit ist waagerechter OKN bei gesunden Erwachsenen symmetrisch " + R("23477776", "Abstract") + "."),
        ("Säuglinge: ", "Bei einäugiger Prüfung gibt es in den ersten Lebensmonaten eine Asymmetrie. Erst mit etwa 11 Monaten ist der OKN bei allen symmetrisch " + R("15548816", "Results") + "."),
    ],
    [
        ("Allgemein: ", "Eine deutliche und wiederholbare Asymmetrie ist der klinisch wichtigste Befund beim waagerechten OKN. Sie weist auf eine einseitige Störung der Bewegungsverarbeitung oder der Blicksteuerung hin."),
    ],
    [
        ("Prätektum: ", "Nach Schäden an NOT/DTN war der OKN mit langsamer Phase zur geschädigten Seite vermindert. Die Gegenrichtung war unbeeinträchtigt " + R("2323371", "Abstract") + " (Studie an Affen)."),
        ("Großhirnrinde: ", "Nach Schäden im Bereich des frontalen Augenfelds und des hinteren Scheitellappens waren Folgebewegungen zur Seite der Schädigung verlangsamt, "
         "bei optokinetischem Reiz in 8 von 13 Fällen mit Scheitellappen-Läsion " + R("9010000", "Abstract") + "."),
        ("Hirnhälfte: ", "Bei Schäden der ganzen linken Hirnhälfte war der OKN zur Gegenseite vermindert. Bei reinen Scheitellappen-Läsionen fand sich kein Richtungsüberwiegen " + R("11677733", "Abstract") + "."),
        ("Muskeln: ", "Ist die Asymmetrie muskulär bedingt, sind M. rectus lateralis (VI) oder medialis (III) einer Seite betroffen. Mit dem Mittelwert beider Augen lässt sich das nicht trennen."),
    ],
    [
        ("Internukleäre Ophthalmoplegie: ", "Bei Läsion des medialen Längsbündels ist die Adduktion des Auges auf der Seite der Läsion gestört " + R("25145891", "Abstract") + ". "
         "Folgebewegungen und OKN wurden dabei untersucht " + R("666605", "Abstract") + ". Die App kann das nicht erkennen, weil sie beide Augen mittelt."),
    ],
    [
        ("Gleiche Person: ", "Eine Asymmetrie, die neu auftritt oder deutlich zunimmt, ist der aussagekräftigste Vergleichsbefund. Sie sollte mit gleichem Aufbau wiederholt bestätigt werden."),
        ("Verschiedene Personen: ", "Wie groß die normale Streuung der waagerechten Asymmetrie bei der App ist, muss erst eine Sammlung von Messungen Gesunder zeigen."),
    ],
)

measure(
    "4.7 Asymmetrie oben/unten",
    "(Gain oben − Gain unten) ÷ Mittelwert beider × 100 %. Positiv: Nach oben laufende Streifen werden besser verfolgt.",
    "wie Abschnitt 4.6, mit den Richtungen ↑ und ↓.",
    [
        ("Studienlage uneinheitlich: ", "Mehrere Studien fanden einen Vorteil nach oben, im Mittel etwa 0,15 Gain " + R("3384058", "Abstract") + R("12939299", "Results")[0:0] + "; "
         "ebenso " + R("12939299", "Results") + " und bei 30–70 °/s " + R("2792261", "Abstract") + ". "
         "Andere fanden in der Gruppe keine senkrechte Asymmetrie " + R("18235002", "Results") + "."),
        ("Persönliche Eigenart: ", "Senkrechte Asymmetrien sind individuell verschieden, bei derselben Person aber gleichbleibend " + R("18235002", "Conclusions") + "."),
        ("Reizfeld: ", "Die Oben-unten-Asymmetrie hängt vom Rand des Gesichtsfeldes ab. Bei rein zentralem Reiz verschwand sie " + R("2792261", "Abstract") + ". "
         "Auf dem iPad ist sie deshalb vermutlich kleiner als im Labor."),
    ],
    [
        ("Allgemein: ", "Wegen der großen individuellen Unterschiede eignet sich der senkrechte Wert vor allem für den Verlauf bei derselben Person, weniger für den Vergleich zwischen Personen."),
    ],
    [
        ("Senkrechte Blickzentren: ", "riMLF, Nucleus interstitialis Cajal und Commissura posterior im Mittelhirn " + R("25145891", "Abstract") + R("16636788", "Abstract") + "."),
        ("Muskeln: ", "nach oben M. rectus superior + M. obliquus inferior (III), nach unten M. rectus inferior (III) + M. obliquus superior (IV)."),
    ],
    [
        ("Parkinson: ", "Der senkrechte OKN war stark asymmetrisch, mit vermindertem Gain für abwärts laufende Reize, deutlicher als bei Kontrollen " + R("19514011", "Abstract") + "."),
        ("PSP: ", "Der senkrechte Gain war reduziert und die senkrechten Rücksprünge verlangsamt " + R("15377682", "Results") + ". Ein isolierter Ausfall senkrechter Augenbewegungen weist auf das Mittelhirn hin " + R("25145891", "Abstract") + "."),
        ("riMLF: ", "Schon eine einseitige Schädigung des riMLF kann eine Blickparese nach oben oder nach oben und unten verursachen " + R("2303833", "Abstract") + "."),
    ],
    [
        ("Gleiche Person: ", "Weil die persönliche senkrechte Asymmetrie stabil ist " + R("18235002", "Conclusions") + ", ist eine Veränderung gegenüber den eigenen früheren Werten aussagekräftiger als der Wert selbst."),
        ("Verschiedene Personen: ", "Unterschiede zwischen Personen sind beim senkrechten Wert häufig normal."),
    ],
)

H("4.8 Technische Kennwerte", level=2)
P("Diese Werte beschreiben die Messqualität, nicht die Hirnfunktion:")
T(["Kennwert", "Bedeutung", "Was tun bei schlechten Werten"], [
    ["Augen erkannt (%)", "Anteil gültiger Bilder im Block", "Licht, Brille, Kamera auf Augenhöhe"],
    ["Wackeln (%)", "Anteil der Bilder mit Geräte- oder Kopfbewegung", "Ständer verwenden, Kopf abstützen"],
    ["eff. °/s", "Streifengeschwindigkeit, auf den tatsächlichen Abstand umgerechnet", "Weicht sie stark vom Sollwert ab, war der Abstand anders als eingestellt"],
    ["Qualität", "gut / mäßig / ungültig (Abschnitt 2.5)", "Ungültige Blöcke wiederholen"],
    ["Kalibrierung R² / Trennung", "Güte der Umrechnung in Grad", "Bei knappen Werten neu kalibrieren"],
], [3.6, 6.4, 6.6])

H("4.9 Nicht gemessen: Nachnystagmus", level=2)
P("Nach Ende der Streifenbewegung im Dunkeln läuft der Nystagmus kurz weiter (optokinetischer Nachnystagmus, OKAN). Er spiegelt den "
  "Geschwindigkeitsspeicher in Vestibulariskernen und Kleinhirn wider " + R("3346723", "Abstract") + ". Nach Schäden an NOT/DTN war er "
  "vermindert oder fehlte " + R("2323371", "Abstract") + ". Bei MS war seine Zeitkonstante verkürzt " + R("11388418", "Abstract") + ". "
  "Bei Patienten mit anhaltenden Beschwerden nach Gehirnerschütterung war sie verlängert " + R("33367947", "Abstract") + ". "
  "Die App misst den OKAN noch nicht. Er wäre eine mögliche Erweiterung: Bildschirm am Blockende schwarz schalten und weitermessen.")

# ---------------------------------------------------------------- 5
H("5 Vergleich von zwei oder mehr Messungen", level=1)
H("5.1 Was die App anzeigt", level=2)
B("Die älteste ausgewählte Messung ist der Bezug.", "Bezug: ")
B("Gain gesamt, waagerecht und senkrecht, Asymmetrie rechts/links und oben/unten, Schlagfrequenz und Amplitude, jeweils mit Abweichung zum Bezug "
  "(absolut, beim Gain auch in %).", "Tabelle „Zusammenfassung“: ")
B("Gain je Richtung und Geschwindigkeit mit Abweichung.", "Tabelle „Gain je Richtung und Geschwindigkeit“: ")
B("Vier Diagramme, eines pro Richtung, mit einer Linie pro Messung über 10, 20 und 30 °/s.", "Diagramme „Gain je Richtung“: ")
B("Balken für waagerechte und senkrechte Asymmetrie, ein Balken pro Messung.", "Diagramm „Asymmetrie“: ")

H("5.2 Wann ist ein Unterschied bedeutsam?", level=2)
P("Für die App gibt es noch keine Studien zur Messgenauigkeit. Bis dahin gilt diese praktische Vorgehensweise. "
  "Sie ist eine Empfehlung und kein Literaturwert:")
N("Pro Person zu Beginn 2–3 Messungen am selben Tag machen. Wie stark diese voneinander abweichen, ist die eigene Messschwankung.", "Ausgangslage: ")
N("Ein späterer Unterschied gilt erst als echt, wenn er deutlich größer als diese Schwankung ist und sich bei einer Wiederholung bestätigt.", "Echte Veränderung: ")
N("Ein Unterschied, der nur eine Richtung betrifft, ist aussagekräftiger als eine gleichmäßige Veränderung aller Richtungen. "
  "Gleichmäßige Veränderungen haben oft äußere Ursachen wie Müdigkeit, Licht oder Abstand.", "Muster vor Einzelwert: ")
N("Qualität und Wackeln der verglichenen Blöcke prüfen. Ein „mäßiger“ Block kann einen Unterschied vortäuschen.", "Qualität: ")

H("5.3 Veränderungsmuster und mögliche Bedeutung", level=2)
P("Die Tabelle zeigt, welche Muster in Studien mit welchen Hirnbereichen in Verbindung gebracht wurden. "
  "Sie ist eine Orientierung und keine Diagnoseanleitung.")
T(["Muster im Vergleich", "Mögliche Bedeutung", "Hirnbereich / Augenmuskel", "Quelle"], [
    ["Gain in allen Richtungen gleichmäßig niedriger", "Aufmerksamkeit, Müdigkeit, Aufbau; Alter; allgemeine Schwäche des Folgesystems",
     "Kleinhirn (Uvula, Pyramis); altersbedingt", R("19567703", "Abstract") + ", " + R("15548816", "Results") + ", " + R("16122840", "Abstract")],
    ["Eine waagerechte Richtung niedriger (neue Asymmetrie R/L)", "einseitige Störung der Bewegungsverarbeitung oder Folgesteuerung",
     "NOT/DTN der Seite, zu der die schwächere langsame Phase zeigt (Tierstudie); frontales Augenfeld / hinterer Scheitellappen; Hirnhälfte",
     R("2323371", "Abstract") + ", " + R("9010000", "Abstract") + ", " + R("11677733", "Abstract")],
    ["Senkrecht nach unten schlechter als früher", "veränderte senkrechte Folgebewegung",
     "Basalganglien-Einfluss (Parkinson-Befund)", R("19514011", "Abstract")],
    ["Senkrechter Gain niedriger, Rücksprünge selten oder schräg", "gestörtes senkrechtes Sakkadensystem",
     "Mittelhirn: riMLF; Muskeln III/IV", R("15377682", "Conclusions") + ", " + R("7066670", "Abstract")],
    ["Waagerechte Rücksprünge fehlen bei erhaltenem Gain", "gestörte waagerechte Sakkadenerzeugung",
     "Brücke: PPRF", R("11960796", "Abstract") + ", " + R("25145891", "Abstract")],
    ["Nur Frequenz/Amplitude anders, Gain gleich", "andere Blickstrategie (Look/Stare), Aufmerksamkeit",
     "kortikale Augenfelder bei Look-OKN", R("15864563", "Abstract") + ", " + R("20403339", "Abstract")],
    ["Gain stärker abfallend bei 30 °/s", "geringere Reserve des schnellen Folgesystems",
     "Brückenkerne, Flocculus, Uvula", R("28646240", "Abstract")],
    ["Gain höher als früher", "z. B. Look- statt Stare-Anweisung; nach Gehirnerschütterung wurde eine höhere visuelle Verstärkung beschrieben (andere Reizform)",
     "subkortikale Blickstabilisierung", R("18235002", "Purpose") + ", " + R("38607620", "Conclusions")],
], [3.6, 4.4, 4.6, 4.0])

H("5.4 Vergleich derselben Person über die Zeit", level=2)
B("Das ist die aussagekräftigste Anwendung. Persönliche Eigenheiten, etwa die senkrechte Asymmetrie, bleiben bei derselben Person gleich " + R("18235002", "Conclusions") + ".", "Stärke: ")
B("Ab etwa 50 Jahren sinkt der Gain langsam " + R("15548816", "Results") + ". Bei Hochbetagten nahm er über Jahre ab, gemeinsam mit Gang und Gleichgewicht " + R("16122840", "Abstract") + ". "
  "Über Jahre hinweg ist eine leichte Abnahme deshalb normal.", "Alterseffekt: ")
B("Gleiches Gerät, gleicher Abstand, gleiche Haltung, gleiche Anweisung, ähnliche Tageszeit, gleiches Licht.", "Voraussetzungen: ")

H("5.5 Vergleich verschiedener Personen", level=2)
B("Nur Personen ähnlichen Alters vergleichen " + R("15548816", "Results") + R("26410670", "Abstract") + ".", "Alter: ")
B("Senkrechte Asymmetrien sind individuell verschieden und gehören zur Person " + R("18235002", "Conclusions") + ". Waagerechte Asymmetrien sind bei Gesunden selten " + R("23477776", "Abstract") + ".", "Asymmetrien: ")
B("Messungen „in der Hand“ nicht direkt mit Messungen „im Ständer“ vergleichen. Die App markiert Handmessungen im Vergleich.", "Haltung: ")

# ---------------------------------------------------------------- 6
H("6 Empfehlungen für gute Messungen", level=1)
for lead, txt in [
    ("Aufbau: ", "iPad quer im Ständer, Kamera auf Augenhöhe, Abstand mit Maßband messen."),
    ("Licht: ", "gleichmäßig von vorne, kein Gegenlicht vom Fenster."),
    ("Brille: ", "spiegelnde Gläser erschweren die Erkennung. Für Vergleiche immer gleich verfahren (mit oder ohne)."),
    ("Anweisung: ", "jedes Mal wörtlich gleich, zum Beispiel: „Schauen Sie entspannt auf die Streifen.“"),
    ("Einstellungen: ", "Streifenbreite, Geschwindigkeiten und Blockdauer für eine Messreihe nie ändern."),
    ("Wiederholung: ", "ungültige Blöcke sofort wiederholen. Zu Beginn mehrere Ausgangsmessungen machen."),
    ("Sichern: ", "nach jeder Sitzung exportieren, zum Beispiel per AirDrop oder in SeaFile."),
]:
    B(txt, lead)

# ---------------------------------------------------------------- 7
H("7 Begriffe", level=1)
T(["Begriff", "Erklärung"], [
    ["OKR / OKN", "optokinetischer Reflex / optokinetischer Nystagmus: Folgebewegung der Augen auf ein bewegtes Muster"],
    ["Langsame Phase", "Folgebewegung in Laufrichtung"],
    ["Schnelle Phase", "Rücksprung (Sakkade) gegen die Laufrichtung"],
    ["Gain", "Augengeschwindigkeit ÷ Streifengeschwindigkeit"],
    ["Look / Stare", "aktives Verfolgen / passives Hinschauen"],
    ["NOT / DTN", "Nucleus tractus optici / dorsaler Terminalkern: Eingangsstation des OKR im Prätektum"],
    ["MT/V5, MST", "Bewegungsareale der Großhirnrinde"],
    ["PPRF", "paramediane pontine Formatio reticularis: Zentrum für waagerechte Sakkaden (Brücke)"],
    ["riMLF", "rostraler interstitieller Kern des medialen Längsbündels: Zentrum für senkrechte Sakkaden (Mittelhirn)"],
    ["MLF", "mediales Längsbündel: verbindet Abducens- und Okulomotoriuskern"],
    ["Flocculus, Uvula, Vermis", "Teile des Kleinhirns, die Augenbewegungen feinsteuern"],
    ["Velocity storage", "Geschwindigkeitsspeicher in Vestibulariskernen und Kleinhirn"],
    ["OKAN", "optokinetischer Nachnystagmus: Weiterlaufen des Nystagmus im Dunkeln"],
    ["PSP", "progressive supranukleäre Blickparese"],
    ["INO", "internukleäre Ophthalmoplegie: Adduktionsschwäche bei MLF-Läsion"],
], [3.8, 12.8])

# ---------------------------------------------------------------- 8
H("8 Literatur", level=1)
P("Alle Arbeiten sind in PubMed gelistet (https://pubmed.ncbi.nlm.nih.gov/<PMID>). Die Abschnittsangaben im Text beziehen sich auf das PubMed-Abstract.",
  size=9.5, color=INK2)
for i, (pmid, cit, doi, pmc) in enumerate(REFS, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"[{i}]\t"); r.bold = True; r.font.size = Pt(9.5)
    r = p.add_run(cit + " "); r.font.size = Pt(9.5)
    extra = f"PMID {pmid} · doi:{doi}" + (f" · {pmc} (frei verfügbar)" if pmc else "")
    r = p.add_run(extra); r.font.size = Pt(9); r.font.color.rgb = INK2

footer()
doc.save(OUT)
print("gespeichert:", OUT, "Literatur:", len(REFS))
